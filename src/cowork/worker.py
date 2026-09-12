"""Deterministic Cowork file worker. Paths are supplied only by LAIN's owned wrapper."""

import csv
import io
import json
import os
import re
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime
from html import escape
from pathlib import Path
from xml.etree import ElementTree


class CoworkError(Exception):
    def __init__(self, kind, why):
        super().__init__(why)
        self.kind = kind


def emit(**values):
    sys.stdout.write(json.dumps(values, separators=(",", ":"), default=str))


def bounded_archive(source):
    try:
        with zipfile.ZipFile(source) as archive:
            entries = archive.infolist()
            if len(entries) > 10_000 or sum(item.file_size for item in entries) > 64 * 1024 * 1024 \
                    or any(item.file_size > 32 * 1024 * 1024 for item in entries):
                raise CoworkError("UNSUPPORTED", "compressed office document expands beyond Cowork limits")
    except zipfile.BadZipFile:
        raise CoworkError("FAILED", "office document is not a valid ZIP package")


def selected_sheets(workbook, names):
    if not names:
        return list(workbook.worksheets)
    missing = [name for name in names if name not in workbook.sheetnames]
    if missing:
        raise CoworkError("FAILED", "unknown worksheet: " + ", ".join(missing[:3]))
    return [workbook[name] for name in names]


def header_map(sheet, row):
    return {str(cell.value).strip(): cell.column for cell in sheet[row] if cell.value is not None}


def column_numbers(sheet, operation):
    names = operation.get("columns") or []
    if not names:
        return list(range(1, sheet.max_column + 1))
    headers = header_map(sheet, max(1, int(operation.get("header_row") or 1)))
    missing = [str(name) for name in names if str(name) not in headers]
    if missing:
        raise CoworkError("FAILED", "unknown column: " + ", ".join(missing[:5]))
    return [headers[str(name)] for name in names]


def as_date(value):
    if not isinstance(value, str) or not value.strip():
        return None
    text = value.strip()
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).replace(tzinfo=None)
    except ValueError:
        pass
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%Y/%m/%d", "%d-%b-%Y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def numeric_summary(rows, headers):
    result = {}
    for index, header in enumerate(headers[:50]):
        values = []
        for row in rows[:5000]:
            if index >= len(row) or isinstance(row[index], bool):
                continue
            try:
                values.append(float(row[index]))
            except (TypeError, ValueError):
                pass
        if values:
            result[str(header or f"column_{index + 1}")] = {
                "count": len(values), "min": min(values), "max": max(values), "sum": sum(values),
                "mean": sum(values) / len(values),
            }
    return result


def fit_columns(sheet):
    for column in range(1, min(sheet.max_column, 200) + 1):
        size = 0
        for row in range(1, min(sheet.max_row, 5000) + 1):
            value = sheet.cell(row, column).value
            size = max(size, len(str(value)) if value is not None else 0)
        sheet.column_dimensions[sheet.cell(1, column).column_letter].width = min(50, max(9, size + 2))


def spreadsheet(payload):
    source = Path(payload.get("input") or "")
    action = payload.get("action")
    if action == "create":
        try:
            import openpyxl
        except ImportError:
            raise CoworkError("UNSUPPORTED", "spreadsheet creation needs Python openpyxl")
        book = openpyxl.Workbook()
        book.remove(book.active)
        total = 0
        for index, spec in enumerate((payload.get("sheets_data") or [])[:20]):
            name = str(spec.get("name") or f"Sheet{index + 1}")[:31]
            sheet = book.create_sheet(name)
            for row in (spec.get("rows") or [])[:5000]:
                values = list(row)[:200]
                total += len(values)
                if total > 100000:
                    raise CoworkError("UNSUPPORTED", "workbook creation exceeds 100,000 cells")
                sheet.append(values)
        if not book.worksheets:
            book.create_sheet("Sheet1")
        output = Path(payload["output"])
        book.save(output)
        openpyxl.load_workbook(output, read_only=True, data_only=False).close()
        return {"ok": True, "changed": total, "facts": {"format": "xlsx", "sheets": len(book.worksheets)}}
    if source.suffix.lower() == ".xls":
        raise CoworkError("UNSUPPORTED", "legacy .xls needs conversion to .xlsx before Cowork can edit it")
    if source.suffix.lower() == ".csv":
        with source.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))
        headers = rows[0][:50] if rows else []
        if action == "inspect":
            return {"ok": True, "facts": {"format": "csv", "rows": len(rows), "columns": max([len(r) for r in rows] or [0]),
                "headers": headers, "numeric": numeric_summary(rows[1:], headers)}}
        changed = 0
        for operation in payload.get("operations") or []:
            op = operation.get("op")
            if op == "trim_text":
                for row in rows:
                    for index, value in enumerate(row):
                        clean = value.strip()
                        if clean != value:
                            row[index] = clean
                            changed += 1
            elif op == "remove_blank_rows":
                before = len(rows)
                rows = [row for i, row in enumerate(rows) if i == 0 or any(str(v).strip() for v in row)]
                changed += before - len(rows)
            elif op == "deduplicate":
                seen, kept = set(), []
                for i, row in enumerate(rows):
                    key = tuple(row)
                    if i and key in seen:
                        changed += 1
                    else:
                        kept.append(row)
                        if i:
                            seen.add(key)
                rows = kept
            elif op == "sort_rows":
                columns = operation.get("columns") or []
                indexes = [headers.index(name) for name in columns if name in headers]
                if len(indexes) != len(columns):
                    raise CoworkError("FAILED", "sort column was not found")
                rows = rows[:1] + sorted(rows[1:], key=lambda row: tuple(row[i] if i < len(row) else "" for i in indexes),
                    reverse=bool(operation.get("descending")))
                changed += 1
            else:
                raise CoworkError("UNSUPPORTED", f"{op} is not supported for CSV")
        with Path(payload["output"]).open("w", encoding="utf-8-sig", newline="") as handle:
            csv.writer(handle).writerows(rows)
        return {"ok": True, "changed": changed, "facts": {"format": "csv", "rows": len(rows), "columns": max([len(r) for r in rows] or [0])}}

    try:
        import openpyxl
        from openpyxl.chart import BarChart, LineChart, PieChart, Reference
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError:
        raise CoworkError("UNSUPPORTED", "xlsx/xlsm support needs Python openpyxl")
    bounded_archive(source)
    keep_vba = source.suffix.lower() == ".xlsm"
    book = openpyxl.load_workbook(source, data_only=False, keep_vba=keep_vba, keep_links=True)
    if action == "inspect":
        facts = {"format": source.suffix.lower().lstrip("."), "sheets": []}
        for sheet in book.worksheets:
            headers = [sheet.cell(1, c).value for c in range(1, min(sheet.max_column, 50) + 1)]
            values = [[sheet.cell(r, c).value for c in range(1, len(headers) + 1)] for r in range(2, min(sheet.max_row, 5001) + 1)]
            facts["sheets"].append({"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column,
                "headers": headers, "numeric": numeric_summary(values, headers)})
        return {"ok": True, "facts": facts}
    changed = 0
    sheets = selected_sheets(book, payload.get("sheets") or [])
    for operation in payload.get("operations") or []:
        op = operation.get("op")
        for sheet in sheets:
            header_row = max(1, int(operation.get("header_row") or 1))
            if op == "trim_text":
                for row in sheet.iter_rows(min_row=header_row + 1):
                    for cell in row:
                        if isinstance(cell.value, str) and not cell.value.startswith("="):
                            clean = cell.value.strip()
                            if clean != cell.value:
                                cell.value = clean
                                changed += 1
            elif op == "remove_blank_rows":
                for row in range(sheet.max_row, header_row, -1):
                    if not any(sheet.cell(row, c).value not in (None, "") for c in range(1, sheet.max_column + 1)):
                        sheet.delete_rows(row)
                        changed += 1
            elif op == "deduplicate":
                columns = column_numbers(sheet, operation)
                seen, duplicate_rows = set(), []
                for row in range(header_row + 1, sheet.max_row + 1):
                    key = tuple(sheet.cell(row, column).value for column in columns)
                    if key in seen:
                        duplicate_rows.append(row)
                    else:
                        seen.add(key)
                for row in reversed(duplicate_rows):
                    sheet.delete_rows(row)
                    changed += 1
            elif op == "normalize_dates":
                for column in column_numbers(sheet, operation):
                    for row in range(header_row + 1, sheet.max_row + 1):
                        cell = sheet.cell(row, column)
                        parsed = as_date(cell.value)
                        if parsed:
                            cell.value, cell.number_format = parsed, "yyyy-mm-dd"
                            changed += 1
            elif op == "sort_rows":
                columns = column_numbers(sheet, operation)
                body = list(sheet.iter_rows(min_row=header_row + 1, values_only=True))
                body.sort(key=lambda row: tuple("" if row[c - 1] is None else str(row[c - 1]) for c in columns),
                    reverse=bool(operation.get("descending")))
                for rindex, values in enumerate(body, header_row + 1):
                    for cindex, value in enumerate(values, 1):
                        sheet.cell(rindex, cindex).value = value
                changed += 1
            elif op == "format_table":
                fill, edge = PatternFill("solid", fgColor="1F4E78"), Side(style="thin", color="D9E2F3")
                for cell in sheet[header_row]:
                    if cell.value is not None:
                        cell.fill = fill
                        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                        cell.border = Border(bottom=edge)
                for row in sheet.iter_rows(min_row=header_row + 1):
                    for cell in row:
                        cell.font, cell.alignment = Font(name="Arial", size=10), Alignment(vertical="center")
                sheet.freeze_panes = sheet.cell(header_row + 1, 1)
                sheet.auto_filter.ref, sheet.sheet_view.showGridLines = sheet.dimensions, False
                fit_columns(sheet)
                changed += 1
            elif op == "autofit":
                fit_columns(sheet)
                changed += 1
            elif op == "chart":
                chart_type = operation.get("chart_type") or "bar"
                chart = {"bar": BarChart, "line": LineChart, "pie": PieChart}.get(chart_type)
                if not chart:
                    raise CoworkError("UNSUPPORTED", "chart type must be bar, line or pie")
                columns = column_numbers(sheet, operation)
                if not columns:
                    raise CoworkError("FAILED", "chart needs a data column")
                drawing = chart()
                drawing.title = str(operation.get("title") or "Summary")[:120]
                drawing.add_data(Reference(sheet, min_col=columns[0], min_row=header_row, max_row=sheet.max_row), titles_from_data=True)
                if len(columns) > 1:
                    drawing.set_categories(Reference(sheet, min_col=columns[1], min_row=header_row + 1, max_row=sheet.max_row))
                sheet.add_chart(drawing, str(operation.get("anchor") or "H2")[:10])
                changed += 1
            else:
                raise CoworkError("UNSUPPORTED", f"unknown spreadsheet operation: {op}")
    output = Path(payload["output"])
    book.save(output)
    verify = openpyxl.load_workbook(output, read_only=True, data_only=False, keep_vba=keep_vba, keep_links=True)
    facts = {"format": output.suffix.lower().lstrip("."), "sheets": [{"name": s.title, "rows": s.max_row, "columns": s.max_column} for s in verify.worksheets]}
    verify.close()
    return {"ok": True, "changed": changed, "facts": facts}


def image_work(payload):
    try:
        from PIL import Image, ImageEnhance, ImageFilter, ImageOps
    except ImportError:
        raise CoworkError("UNSUPPORTED", "image support needs Python Pillow")
    source = Path(payload["input"])
    current = Image.open(source)
    if current.width * current.height > 40_000_000:
        raise CoworkError("UNSUPPORTED", "input image exceeds 40 megapixels")
    current.load()
    current = ImageOps.exif_transpose(current)
    if payload.get("action") == "inspect":
        return {"ok": True, "facts": {"format": current.format or source.suffix.lstrip("."), "width": current.width, "height": current.height, "mode": current.mode}}
    for operation in payload.get("operations") or []:
        op = operation.get("op")
        if op in ("resize", "upscale"):
            if op == "upscale":
                factor = max(1.0, min(4.0, float(operation.get("factor") or 2)))
                width, height = round(current.width * factor), round(current.height * factor)
            else:
                width, height = int(operation.get("width") or 0), int(operation.get("height") or 0)
                if width <= 0 and height <= 0:
                    raise CoworkError("FAILED", "resize needs a positive width or height")
                if width <= 0:
                    width = max(1, round(current.width * height / current.height))
                if height <= 0:
                    height = max(1, round(current.height * width / current.width))
            if width * height > 40_000_000:
                raise CoworkError("UNSUPPORTED", "requested image exceeds 40 megapixels")
            current = current.resize((width, height), Image.Resampling.LANCZOS)
        elif op == "crop":
            x, y = int(operation.get("x") or 0), int(operation.get("y") or 0)
            width, height = int(operation.get("width") or 0), int(operation.get("height") or 0)
            if min(x, y) < 0 or width <= 0 or height <= 0 or x + width > current.width or y + height > current.height:
                raise CoworkError("FAILED", "crop rectangle is outside the image")
            current = current.crop((x, y, x + width, y + height))
        elif op == "rotate":
            current = current.rotate(float(operation.get("degrees") or 0), expand=bool(operation.get("expand", True)), resample=Image.Resampling.BICUBIC)
        elif op == "sharpen":
            current = ImageEnhance.Sharpness(current).enhance(max(0.0, min(5.0, float(operation.get("factor") or 1.5))))
        elif op == "denoise":
            radius = max(1, min(5, int(operation.get("radius") or 1)))
            current = current.filter(ImageFilter.MedianFilter(size=radius * 2 + 1))
        elif op == "grayscale":
            current = ImageOps.grayscale(current)
        elif op == "flatten_background":
            rgba = current.convert("RGBA")
            background = Image.new("RGBA", rgba.size, str(operation.get("color") or "white"))
            background.alpha_composite(rgba)
            current = background.convert("RGB")
        elif op == "remove_background":
            runner = payload.get("background_runner") or {}
            if runner.get("program") and runner.get("script"):
                source_temp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
                result_temp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
                source_temp.close(); result_temp.close()
                try:
                    current.save(source_temp.name, format="PNG")
                    done = subprocess.run([runner["program"], runner["script"], source_temp.name, result_temp.name],
                        capture_output=True, text=True, timeout=180, check=False)
                    if done.returncode or not Path(result_temp.name).is_file() or Path(result_temp.name).stat().st_size == 0:
                        raise CoworkError("FAILED", "the configured background-removal provider failed")
                    current = Image.open(result_temp.name)
                    current.load()
                except subprocess.TimeoutExpired:
                    raise CoworkError("FAILED", "the configured background-removal provider timed out")
                finally:
                    for temporary in (source_temp.name, result_temp.name):
                        try: os.unlink(temporary)
                        except OSError: pass
            else:
                try:
                    from rembg import remove
                except ImportError:
                    raise CoworkError("UNSUPPORTED", "background removal needs a configured rembg/BiRefNet provider")
                removed = remove(current)
                if isinstance(removed, bytes):
                    current = Image.open(io.BytesIO(removed)); current.load()
                else:
                    current = removed
        else:
            raise CoworkError("UNSUPPORTED", f"unknown image operation: {op}")
    output = Path(payload["output"])
    fmt = str(payload.get("format") or output.suffix.lstrip(".") or "png").upper().replace("JPG", "JPEG")
    if fmt == "JPEG" and current.mode not in ("RGB", "L"):
        current = current.convert("RGB")
    current.save(output, format=fmt, quality=max(1, min(100, int(payload.get("quality") or 92))))
    check = Image.open(output); check.verify()
    return {"ok": True, "changed": len(payload.get("operations") or []), "facts": {"format": fmt.lower(), "width": current.width, "height": current.height, "mode": current.mode}}


def docx_text(source):
    bounded_archive(source)
    with zipfile.ZipFile(source) as archive:
        xml = archive.read("word/document.xml")
    root = ElementTree.fromstring(xml)
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    paragraphs = []
    for paragraph in root.iter(ns + "p"):
        text = "".join(node.text or "" for node in paragraph.iter(ns + "t"))
        if text:
            paragraphs.append(text)
    return "\n".join(paragraphs)


def write_docx(output, title, paragraphs):
    def para(text, style=""):
        props = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ""
        return f'<w:p>{props}<w:r><w:t xml:space="preserve">{escape(str(text))}</w:t></w:r></w:p>'
    body = (para(title, "Title") if title else "") + "".join(para(value) for value in paragraphs)
    document = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>' + body + '<w:sectPr><w:pgSz w:w="12240" w:h="15840"/><w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440"/></w:sectPr></w:body></w:document>'
    content_types = '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>'
    rels = '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>'
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", rels)
        archive.writestr("word/document.xml", document)


def write_pdf(output, title, paragraphs):
    lines = ([title] if title else []) + [line for text in paragraphs for line in str(text).splitlines()]
    pages = [lines[i:i + 48] for i in range(0, max(1, len(lines)), 48)] or [[]]
    objects = []
    def add(body):
        objects.append(body)
        return len(objects)
    catalog = add("")
    pages_id = add("")
    font_id = add("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    page_ids = []
    for page in pages:
        commands = ["BT /F1 11 Tf 72 750 Td 14 TL"]
        for line in page:
            clean = str(line).encode("latin-1", "replace").decode("latin-1").replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            commands.append(f"({clean[:120]}) Tj T*")
        commands.append("ET")
        stream = "\n".join(commands).encode("latin-1")
        stream_id = add(f"<< /Length {len(stream)} >>\nstream\n" + stream.decode("latin-1") + "\nendstream")
        page_ids.append(add(f"<< /Type /Page /Parent {pages_id} 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {stream_id} 0 R >>"))
    objects[catalog - 1] = f"<< /Type /Catalog /Pages {pages_id} 0 R >>"
    objects[pages_id - 1] = f"<< /Type /Pages /Kids [{' '.join(f'{item} 0 R' for item in page_ids)}] /Count {len(page_ids)} >>"
    data = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for index, body in enumerate(objects, 1):
        offsets.append(len(data)); data.extend(f"{index} 0 obj\n{body}\nendobj\n".encode("latin-1"))
    xref = len(data)
    data.extend(f"xref\n0 {len(objects)+1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]: data.extend(f"{offset:010d} 00000 n \n".encode())
    data.extend(f"trailer << /Size {len(objects)+1} /Root {catalog} 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    Path(output).write_bytes(data)


def extract_document(source):
    ext = source.suffix.lower()
    if ext in (".txt", ".md", ".csv"):
        return source.read_text("utf-8-sig", errors="replace"), None
    if ext == ".docx":
        return docx_text(source), None
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            raise CoworkError("UNSUPPORTED", "PDF extraction needs Python pypdf")
        reader = PdfReader(source)
        return "\n\n".join(page.extract_text() or "" for page in reader.pages), len(reader.pages)
    raise CoworkError("UNSUPPORTED", "document must be txt, md, csv, docx or pdf")


def document(payload):
    action, output_format = payload.get("action"), str(payload.get("format") or "docx").lower()
    if action == "create":
        title = str(payload.get("title") or "")[:300]
        paragraphs = [str(value)[:10000] for value in (payload.get("paragraphs") or [])[:500]]
        if output_format == "docx": write_docx(payload["output"], title, paragraphs)
        elif output_format == "pdf": write_pdf(payload["output"], title, paragraphs)
        else: raise CoworkError("UNSUPPORTED", "document output must be docx or pdf")
        return {"ok": True, "changed": len(paragraphs), "facts": {"format": output_format, "paragraphs": len(paragraphs)}}
    source = Path(payload["input"])
    text, pages = extract_document(source)
    facts = {"format": source.suffix.lstrip(".").lower(), "characters": len(text), "words": len(re.findall(r"\S+", text)), "pages": pages,
        "text": text[:60000], "truncated": len(text) > 60000}
    if action == "inspect":
        return {"ok": True, "facts": facts}
    for operation in payload.get("operations") or []:
        op = operation.get("op")
        if op == "replace_text":
            old, new = str(operation.get("find") or ""), str(operation.get("replace") or "")
            if not old: raise CoworkError("FAILED", "replace_text needs find text")
            text = text.replace(old, new)
        elif op == "append_text": text += "\n" + str(operation.get("text") or "")
        else: raise CoworkError("UNSUPPORTED", f"unknown document operation: {op}")
    paragraphs = [part for part in text.splitlines() if part.strip()]
    if output_format == "docx": write_docx(payload["output"], str(payload.get("title") or source.stem), paragraphs)
    elif output_format == "pdf": write_pdf(payload["output"], str(payload.get("title") or source.stem), paragraphs)
    else: raise CoworkError("UNSUPPORTED", "document output must be docx or pdf")
    return {"ok": True, "changed": len(payload.get("operations") or []), "facts": {"format": output_format, "paragraphs": len(paragraphs)}}


def main():
    try:
        payload = json.loads(sys.stdin.read(512_000))
        if not isinstance(payload, dict) or payload.get("kind") not in ("spreadsheet", "image", "document"):
            raise CoworkError("FAILED", "invalid Cowork worker request")
        if payload.get("action") != "create":
            source = str(payload.get("input") or "")
            if not os.path.isabs(source) or not Path(source).is_file():
                raise CoworkError("FAILED", "owned input is unavailable")
        if payload.get("action") not in ("inspect",):
            output = str(payload.get("output") or "")
            if not os.path.isabs(output):
                raise CoworkError("FAILED", "owned output location is unavailable")
        handler = {"spreadsheet": spreadsheet, "image": image_work, "document": document}[payload["kind"]]
        emit(**handler(payload))
    except CoworkError as error:
        emit(ok=False, **{"class": error.kind}, why=str(error))
    except Exception:
        emit(ok=False, **{"class": "FAILED"}, why="the deterministic worker could not process this file")


if __name__ == "__main__":
    main()
